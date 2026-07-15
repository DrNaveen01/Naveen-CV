"""Build the replaceable Word reference document used by Quarto/Pandoc."""

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "templates" / "word" / "cv-reference.docx"

FONT = "Arial"
INK = RGBColor(22, 33, 43)
MUTED = RGBColor(83, 96, 107)
ACCENT = RGBColor(23, 74, 91)


def set_font(style, size, color=INK, bold=False, italic=False):
    style.font.name = FONT
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.font.bold = bold
    style.font.italic = italic


def set_spacing(style, before=0, after=0, line=1.0):
    paragraph = style.paragraph_format
    paragraph.space_before = Pt(before)
    paragraph.space_after = Pt(after)
    paragraph.line_spacing = line


def set_bottom_border(style, color="174A5B", size="8", space="3"):
    p_pr = style._element.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_paragraph_style(document, name, base="Normal"):
    styles = document.styles
    try:
        return styles[name]
    except KeyError:
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = styles[base]
        return style


def build_reference():
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    normal = document.styles["Normal"]
    set_font(normal, 10.25)
    set_spacing(normal, before=0, after=3, line=1.05)

    for style_name in ("Body Text", "First Paragraph"):
        try:
            style = document.styles[style_name]
        except KeyError:
            continue
        set_font(style, 10.25)
        set_spacing(style, before=0, after=3, line=1.05)

    heading_1 = document.styles["Heading 1"]
    set_font(heading_1, 12, color=ACCENT, bold=True)
    set_spacing(heading_1, before=11, after=5, line=1.0)
    heading_1.paragraph_format.keep_with_next = True
    set_bottom_border(heading_1)

    heading_2 = document.styles["Heading 2"]
    set_font(heading_2, 11, color=ACCENT, bold=True)
    set_spacing(heading_2, before=8, after=4, line=1.0)
    heading_2.paragraph_format.keep_with_next = True

    heading_3 = document.styles["Heading 3"]
    set_font(heading_3, 10.25, color=INK, bold=True)
    set_spacing(heading_3, before=6, after=3, line=1.0)
    heading_3.paragraph_format.keep_with_next = True

    cv_name = add_paragraph_style(document, "CV Name")
    set_font(cv_name, 22, color=ACCENT, bold=True)
    set_spacing(cv_name, before=0, after=3, line=1.0)
    cv_name.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cv_name.paragraph_format.keep_with_next = True

    cv_headline = add_paragraph_style(document, "CV Headline")
    set_font(cv_headline, 10.25, color=INK)
    set_spacing(cv_headline, before=0, after=2, line=1.0)
    cv_headline.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cv_headline.paragraph_format.keep_with_next = True

    cv_contact = add_paragraph_style(document, "CV Contact")
    set_font(cv_contact, 9, color=MUTED)
    set_spacing(cv_contact, before=0, after=7, line=1.0)
    cv_contact.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cv_contact.paragraph_format.keep_with_next = True

    cv_summary = add_paragraph_style(document, "CV Summary")
    set_font(cv_summary, 9.5, color=INK)
    set_spacing(cv_summary, before=0, after=8, line=1.05)
    cv_summary.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    cv_entry = add_paragraph_style(document, "CV Entry")
    set_font(cv_entry, 9.8, color=INK)
    set_spacing(cv_entry, before=3, after=0, line=1.05)
    cv_entry.paragraph_format.keep_with_next = True

    cv_detail = add_paragraph_style(document, "CV Detail")
    set_font(cv_detail, 9.4, color=INK)
    set_spacing(cv_detail, before=0, after=2, line=1.05)

    cv_note = add_paragraph_style(document, "CV Note")
    set_font(cv_note, 8.8, color=MUTED, italic=True)
    set_spacing(cv_note, before=0, after=2, line=1.0)

    for list_style_name in ("List Bullet", "List Number"):
        try:
            list_style = document.styles[list_style_name]
        except KeyError:
            continue
        set_font(list_style, 9.4, color=INK)
        set_spacing(list_style, before=0, after=1, line=1.05)
        list_style.paragraph_format.left_indent = Cm(0.5)
        list_style.paragraph_format.first_line_indent = Cm(-0.25)

    # These sample paragraphs ensure every custom style is retained in the
    # reference document. Pandoc uses the style definitions, not this content.
    samples = [
        ("Name", "CV Name"),
        ("Headline", "CV Headline"),
        ("Contact", "CV Contact"),
        ("Summary", "CV Summary"),
        ("Entry", "CV Entry"),
        ("Detail", "CV Detail"),
        ("Note", "CV Note"),
    ]
    for text, style_name in samples:
        document.add_paragraph(text, style=style_name)

    document.core_properties.title = "Naveen CV Word reference template"
    document.core_properties.subject = "Replaceable Word style layer"
    document.core_properties.author = "Naveen"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_reference()

